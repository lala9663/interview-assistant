"""Generate the synthetic resumes for the file-based leak check (tests/fixtures/resumes).

    uv run python scripts/make_resume_fixtures.py
    uv run python scripts/make_resume_fixtures.py --font /usr/share/fonts/truetype/nanum/NanumGothic.ttf

Twenty fictional applicants in five forms: text PDF 6, DOCX 4, PPTX 3, HTML 4, scanned PDF 3.
Every file is a Korean resume (인적사항 table, 학력·경력 tables, 자기소개서) with its own mix
of traps. While a file is drawn, each identifying value is recorded together with the kind of
block it was placed in; manifest.yaml is written from those records, so the ground truth cannot
drift from the files.

Korean font: PDFs and scan images need a TrueType font with Hangul. It is read from the system
(--font, IA_FIXTURE_FONT, or common install paths) and never committed. PDFs embed the subset
of glyphs they use, as any PDF writer does. Needs the dev dependencies (reportlab, pillow).
"""

from __future__ import annotations

import argparse
import io
import os
import re
import string
import sys
import unicodedata
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tests" / "fixtures" / "resumes"

FONT_CANDIDATES = (
    "C:/Windows/Fonts/malgun.ttf",
    "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
    "/usr/share/fonts/nanum/NanumGothic.ttf",
    "/Library/Fonts/NanumGothic.ttf",
    "/System/Library/Fonts/Supplemental/AppleGothic.ttf",
)
FIXED_TIME = datetime(2026, 10, 1, 9, 0, 0)

MANIFEST_HEADER = """\
# 파일 단위 누출 검수용 가상 이력서 20건. 전부 가상 인물이며 실제 개인정보는 없다.
# 생성: uv run python scripts/make_resume_fixtures.py (이 파일은 손으로 고치지 말 것)
#
# 만든 규칙 (tests/fixtures/deid_samples.yaml과 같다)
# - 휴대전화: 010-0xxx-xxxx, 010-1xxx-xxxx (가입자에게 배정되지 않는 국번대)
# - 이메일: example.com / example.org / example.net (RFC 2606 예약 도메인)
# - 주소: 실재 시·군·구에 지어낸 도로명. 학교·회사: 사전의 공개 기관명 또는 지어낸 이름
# - 증명사진: 실루엣만 그린 자리 표시 이미지
#
# file: 이 폴더 기준 경로. upload_name: 업로드 때 파일명이 디스크 이름과 다를 때만 (예: macOS NFD)
# scanned: 텍스트 레이어 없는 스캔본이면 true. OCR 필요로 분류되어야 하고 누출 검사는 보류한다
# identity: 지원서 양식에서 이미 아는 본인 정보 (KnownIdentity)
# expected: 파일 안의 식별 값 정답. where는 그 값을 넣은 블록 종류
#   (filename, metadata, body, table, header, footer, notes, comment, link, hidden,
#    image = 이미지 속 글자라 OCR 전에는 추출되지 않음)
# keep: 식별 정보 옆의 일반 표현. 치환에 같이 먹히면 과잉 치환으로 실패한다
# images: 넣은 이미지 수 (추출기가 버린 이미지 수는 이 이상이어야 한다. 서식 썸네일 포함)
"""


# --- people ---------------------------------------------------------------------------------


@dataclass(frozen=True)
class Person:
    id: str
    form: str  # pdf, docx, pptx, html, scan
    file: str
    name: str
    english: str
    phone: str
    email: str
    birth: str
    region: str  # 시/도 + 시/군/구
    road: str
    unit: str
    school: str
    school_short: str
    major: str
    company: str  # as written in the application form
    company_short: str
    role: str
    job: str
    period: str
    ref: str
    ref_title: str
    template: int
    hanja: str = ""

    @property
    def surname(self) -> str:
        return self.name[:2] if self.name[:2] in ("남궁", "황보", "제갈", "선우") else self.name[0]

    @property
    def given(self) -> str:
        return self.name[len(self.surname) :]

    @property
    def address(self) -> str:
        return f"{self.region} {self.road}" + (f", {self.unit}" if self.unit else "")

    @property
    def phone_digits(self) -> str:
        return self.phone.replace("-", "")

    def identity(self) -> dict:
        aliases = [self.english] + ([self.hanja] if self.hanja else [])
        return {
            "name": self.name,
            "aliases": aliases,
            "phones": [self.phone],
            "emails": [self.email],
            "birthdate": self.birth,
            "addresses": [self.address],
            "schools": [self.school],
            "companies": [self.company],
        }


P = Person
PEOPLE = [
    P("R01", "pdf", "이력서_권하린_2026.pdf", "권하린", "Harin Kwon", "010-0211-3344", "harin.kwon@example.com", "1996-04-17", "서울특별시 마포구", "솔바람길 27", "101동 1203호", "한빛대학교", "한빛대", "컴퓨터공학", "(주)누리소프트", "누리소프트", "백엔드 개발자", "백엔드 개발", "2020.01 ~ 2025.12", "오태호", "팀장", 0, hanja="權夏潾"),
    P("R02", "pdf", "남궁민재_이력서(최종).pdf", "남궁민재", "Minjae Namgung", "010-1302-4455", "mj.namgung@example.org", "1993-08-02", "경기도 성남시 분당구", "새말로 112", "한울아파트 304동 1102호", "고려대학교", "고려대", "산업공학", "카카오", "카카오", "데이터 분석가", "데이터 분석", "2018.07 ~ 2025.09", "문정훈", "파트장", 1),
    P("R03", "pdf", "Resume_SerinOh.pdf", "오세린", "Serin Oh", "010-0447-1290", "serin.oh@example.net", "1998-12-05", "부산광역시 해운대구", "바다숲길 8", "", "부산대학교", "부산대", "전자공학", "LG전자", "LG전자", "회로 설계 엔지니어", "회로 설계", "2021.02 ~ 2025.10", "김태현", "책임", 2),
    P("R04", "pdf", "백승우-자기소개서-최종_v3.pdf", "백승우", "Seungwoo Baek", "010-1588-0921", "swbaek@example.com", "1991-02-28", "대전광역시 유성구", "별빛누리길 42", "한울아파트 103동 502호", "한국과학기술원", "카이스트", "전산학", "새봄테크", "새봄테크", "머신러닝 엔지니어", "머신러닝", "2017.03 ~ 2025.06", "정민지", "박사", 3),
    P("R05", "pdf", "황보은채 이력서.pdf", "황보은채", "Eunchae Hwangbo", "010-0732-6618", "eunchae.hb@example.org", "1997-06-11", "인천광역시 연수구", "물빛로 56", "", "인하대학교", "인하대", "정보통신공학", "다온랩스", "다온랩스", "프론트엔드 개발자", "프론트엔드 개발", "2020.08 ~ 2025.08", "이서준", "선배", 0),
    P("R06", "pdf", "R06_차준호.pdf", "차준호", "Junho Cha", "010-1914-2273", "junho.cha@example.net", "1994-10-30", "광주광역시 북구", "햇살마루길 19", "", "전남대학교", "전남대", "전자공학", "(주)하늘시스템즈", "하늘시스템즈", "임베디드 개발자", "임베디드 소프트웨어", "2019.01 ~ 2025.07", "박하은", "수석", 1),
    P("R07", "docx", "유다인_이력서.docx", "유다인", "Dain Yoo", "010-0825-7731", "dain.yoo@example.com", "1995-01-23", "서울특별시 송파구", "꽃내음로 33", "204동 801호", "성균관대학교", "성균관대", "소프트웨어학", "네이버", "네이버", "백엔드 개발자", "백엔드 개발", "2019.06 ~ 2025.11", "최유진", "리더", 2),
    P("R08", "docx", "입사지원서_탁지호.docx", "탁지호", "Jiho Tak", "010-1176-0458", "jiho.tak@example.org", "1992-09-14", "경기도 수원시 영통구", "누리별로 88", "", "아주대학교", "아주대", "정보컴퓨터공학", "삼성SDS", "삼성SDS", "클라우드 엔지니어", "클라우드 운영", "2018.02 ~ 2025.05", "한지민", "그룹장", 3),
    P("R09", "docx", "선우하람_Resume.docx", "선우하람", "Ha-ram Sunwoo", "010-0369-8812", "haram.sunwoo@example.net", "1999-03-03", "대구광역시 수성구", "들꽃길 21", "", "경북대학교", "경북대", "통계학", "푸른솔시스템즈", "푸른솔시스템즈", "QA 엔지니어", "품질 보증", "2022.01 ~ 2025.12", "윤서준", "매니저", 0),
    P("R10", "docx", "이력서_문태오.docx", "문태오", "Taeo Moon", "010-1450-6627", "taeo.moon@example.com", "1990-11-19", "울산광역시 남구", "은하수로 74", "", "울산과학기술원", "유니스트", "기계공학", "현대자동차", "현대차", "자율주행 연구원", "자율주행 연구", "2016.03 ~ 2025.09", "강민서", "책임", 1, hanja="文泰旿"),
    P("R11", "pptx", "포트폴리오_진서아.pptx", "진서아", "Seoa Jin", "010-0614-2209", "seoa.jin@example.org", "1996-07-07", "서울특별시 성동구", "바람결길 15", "", "한양대학교", "한양대", "산업디자인학", "쿠팡", "쿠팡", "UX 디자이너", "UX 디자인", "2019.09 ~ 2025.10", "서예준", "디자이너", 2),
    P("R12", "pptx", "MinhyukPyo_portfolio.pptx", "표민혁", "Minhyuk Pyo", "010-1027-5583", "minhyuk.pyo@example.net", "1993-05-25", "경기도 고양시 일산동구", "호수빛로 61", "", "홍익대학교", "홍익대", "게임소프트웨어학", "엔씨소프트", "엔씨소프트", "게임 클라이언트 개발자", "게임 클라이언트 개발", "2017.07 ~ 2025.04", "임수아", "실장", 3),
    P("R13", "pptx", "연가은 경력기술서.pptx", "연가은", "Gaeun Yeon", "010-0958-3346", "gaeun.yeon@example.com", "1998-02-14", "세종특별자치시", "한솔빛로 27", "", "충남대학교", "충남대", "컴퓨터공학", "새솔테크", "새솔테크", "데이터 엔지니어", "데이터 엔지니어링", "2021.03 ~ 2025.12", "오지훈", "팀장", 0),
    P("R14", "html", "육지안_이력서.html", "육지안", "Jian Yuk", "010-1233-9087", "jian.yuk@example.org", "1995-09-09", "서울특별시 관악구", "노을빛길 45", "", "서울대학교", "서울대", "컴퓨터공학", "비바리퍼블리카", "토스", "보안 엔지니어", "정보 보안", "2019.02 ~ 2025.08", "남도윤", "팀장", 1),
    P("R15", "html", "jegal_yunho_resume.html", "제갈윤호", "Yunho Jegal", "010-0541-7720", "yunho.jegal@example.net", "1992-04-04", "강원특별자치도 춘천시", "소양숲길 12", "", "강원대학교", "강원대", "컴퓨터공학", "다온랩스", "다온랩스", "인프라 엔지니어", "인프라 운영", "2017.01 ~ 2025.03", "장예은", "선임", 2),
    P("R16", "html", "이력서-마수빈.html", "마수빈", "Subin Ma", "010-1662-3018", "subin.ma@example.com", "1997-11-28", "경상남도 창원시 성산구", "별무리로 9", "", "창원대학교", "창원대", "제어계측공학", "한빛테크", "한빛테크", "제어 소프트웨어 개발자", "제어 소프트웨어", "2021.01 ~ 2025.09", "홍서윤", "매니저", 3),
    P("R17", "html", "하예린 이력서.htm", "하예린", "Yerin Ha", "010-0872-4451", "yerin.ha@example.org", "1994-03-21", "제주특별자치도 제주시", "귤빛로 30", "", "제주대학교", "제주대", "경영학", "카카오", "카카오", "서비스 기획자", "서비스 기획", "2018.03 ~ 2025.07", "최다온", "리더", 0),
    P("R18", "scan", "설도현_이력서_스캔.pdf", "설도현", "Dohyun Seol", "010-1381-0274", "dohyun.seol@example.net", "1993-01-15", "서울특별시 노원구", "초록숲길 5", "", "광운대학교", "광운대", "전자공학", "누리소프트", "누리소프트", "펌웨어 개발자", "펌웨어 개발", "2018.01 ~ 2025.06", "오태호", "팀장", 1),
    P("R19", "scan", "봉채원_scan.pdf", "봉채원", "Chaewon Bong", "010-0296-5530", "chaewon.bong@example.com", "1999-08-08", "경기도 용인시 수지구", "은빛로 18", "", "단국대학교", "단국대", "웹디자인학", "새봄테크", "새봄테크", "웹 퍼블리셔", "웹 퍼블리싱", "2022.02 ~ 2025.10", "정민지", "실장", 2),
    P("R20", "scan", "금태민_지원서.pdf", "금태민", "Taemin Geum", "010-1749-3361", "taemin.geum@example.org", "1991-12-12", "충청북도 청주시 흥덕구", "오송빛로 40", "", "충북대학교", "충북대", "생물정보학", "하늘시스템즈", "하늘시스템즈", "바이오 데이터 분석가", "바이오 데이터", "2016.08 ~ 2025.02", "박하은", "수석", 3),
]


# --- ground truth ---------------------------------------------------------------------------


@dataclass
class Truth:
    """Records each identifying value with the block kinds it was placed in."""

    where: dict[tuple[str, str], list[str]] = field(default_factory=dict)
    keep: list[str] = field(default_factory=list)
    traps: list[str] = field(default_factory=list)
    images: int = 0

    def put(self, kind: str, category: str, value: str) -> str:
        kinds = self.where.setdefault((category, value), [])
        if kind not in kinds:
            kinds.append(kind)
        return value

    def at(self, kind: str) -> Callable[[str, str], str]:
        return lambda category, value: self.put(kind, category, value)

    def trap(self, *names: str) -> None:
        self.traps.extend(n for n in names if n not in self.traps)


# Placeholders in the 자기소개서 templates. Identifying ones carry their category.
def fields(p: Person) -> dict[str, tuple[str | None, str]]:
    return {
        "name": ("candidate", p.name),
        "given": ("candidate", p.given),
        "english": ("candidate", p.english),
        "school": ("school", p.school),
        "school_short": ("school", p.school_short),
        "company": ("company", p.company_short),
        "ref": ("person", p.ref),
        "email": ("email", p.email),
        "phone": ("phone", p.phone),
        "ref_title": (None, p.ref_title),
        "major": (None, p.major),
        "role": (None, p.role),
        "job": (None, p.job),
    }


TEMPLATES = [
    (
        [
            "[지원 동기] 안녕하십니까. {job} 직무에 지원한 {name}입니다. {school_short}에서 {major}을 전공하며 큰 시스템이 실패 없이 돌아가게 만드는 일에 흥미를 느꼈습니다.",
            "[직무 역량] {company}에서 {role} 업무를 맡아 장애 대응 절차를 다시 설계했고, 배포 실패율을 절반으로 줄였습니다. 당시 {ref} {ref_title}님께서 코드 리뷰 문화를 함께 만들어 주셨습니다.",
            "[협업 경험] 동료들은 저를 '{given} 님'이라고 부르며 어려운 일을 먼저 맡기곤 했습니다. 문제를 숫자로 설명하고 합의를 이끄는 것이 제 강점입니다.",
            "[입사 후 포부] 입사 후에는 동료 개발자들이 믿고 쓰는 플랫폼을 만들고 싶습니다. 궁금하신 점은 {email}로 연락 주시면 감사하겠습니다.",
        ],
        ["직무에 지원한", "전공하며", "업무를 맡아", "입사 후 포부"],
    ),
    (
        [
            "[성장 과정] 저는 {name}입니다. 어릴 때부터 고장 난 물건을 분해하고 다시 조립하는 것을 좋아했고, 그 호기심이 {school_short} {major} 진학으로 이어졌습니다.",
            "[경력 사항] {company} 재직 중 {role} 업무를 하며 신규 서비스 세 개를 출시했습니다. {ref} {ref_title}님과 함께한 프로젝트에서 일정과 품질을 함께 지키는 법을 배웠습니다.",
            "[강점] 저는 기록하는 습관이 있습니다. 회의마다 결정 사항을 문서로 남겨 팀이 같은 실수를 반복하지 않도록 했습니다.",
            "[포부] {job} 분야에서 오래 쓰이는 제품을 만들고 싶습니다. 영문 이름 {english}로 공개 저장소에 작업을 올려 두었습니다.",
        ],
        ["고장 난 물건을", "재직 중", "기록하는 습관", "오래 쓰이는 제품"],
    ),
    (
        [
            "[지원 동기] {company}에서 쌓은 {role} 경험을 바탕으로 더 큰 문제를 풀고 싶어 {job} 직무에 지원했습니다.",
            "[학업] {school}에서 {major}을 공부하며 데이터로 의사결정을 설명하는 방법을 익혔습니다.",
            "[협업] {ref} {ref_title}님과 진행한 개선 과제에서 처음으로 팀 전체의 목표를 숫자로 정리해 보았습니다. 이후 모든 일을 측정 가능한 목표로 바꾸어 생각합니다.",
            "[포부] 저 {name}, 입사 후 3년 안에 팀이 믿고 맡기는 사람이 되겠습니다. 연락처는 {phone}입니다.",
        ],
        ["경험을 바탕으로", "의사결정을 설명하는", "측정 가능한 목표로", "믿고 맡기는 사람"],
    ),
    (
        [
            "[자기소개] {name}({english})입니다. {school_short} 졸업 후 {company}에서 일하며 사용자 문제를 가장 먼저 듣는 사람이 되려고 노력했습니다.",
            "[대표 프로젝트] {role} 업무를 하면서 고객 문의를 분류하는 내부 도구를 만들었고, 처리 시간을 40% 줄였습니다. {ref} {ref_title}님이 이 도구를 다른 팀에도 소개해 주셨습니다.",
            "[일하는 방식] 작게 만들고 빨리 확인합니다. 실패한 시도도 기록해 두고 다음 결정의 근거로 씁니다.",
            "[포부] {job} 직무에서 현장의 목소리를 제품으로 연결하겠습니다.",
        ],
        ["사용자 문제를", "고객 문의를 분류하는", "작게 만들고 빨리 확인합니다", "현장의 목소리를"],
    ),
]


def intro(p: Person, truth: Truth, kind: str) -> list[str]:
    """The 자기소개서 paragraphs, recording every identifying placeholder used."""
    paragraphs, keep = TEMPLATES[p.template]
    values = fields(p)
    out = []
    for template in paragraphs:
        for _, name, _, _ in string.Formatter().parse(template):
            if name and values[name][0]:
                truth.put(kind, values[name][0], values[name][1])
        out.append(template.format(**{k: v for k, (_, v) in values.items()}))
    # Job titles sit right next to names and must survive ("인프라 엔지니어" is not a person).
    for phrase in (*keep, p.role, p.job):
        filled = phrase.format(**{k: v for k, (_, v) in values.items()})
        if filled not in truth.keep:
            truth.keep.append(filled)
    return out


# --- value formats --------------------------------------------------------------------------


def phone_as(p: Person, style: str) -> str:
    a, b, c = p.phone.split("-")
    return {
        "dash": p.phone,
        "dots": f"{a}.{b}.{c}",
        "plain": p.phone_digits,
        "spaces": f"{a} {b} {c}",
        "intl": f"+82-{a[1:]}-{b}-{c}",
        "intl0": f"+82 (0){a[1:]} {b} {c}",
        "paren": f"({a}) {b}-{c}",
        "fullwidth": "".join(chr(ord(ch) + 0xFEE0) for ch in p.phone),
    }[style]


def birth_as(p: Person, style: str) -> str:
    y, m, d = p.birth.split("-")
    return {
        "dots": f"{y}.{m}.{d}",
        "korean": f"{y}년 {int(m)}월 {int(d)}일",
        "short": f"{y[2:]}.{m}.{d}",
        "slash": f"{y}/{m}/{d}",
        "iso": p.birth,
    }[style]


def address_pieces(p: Person, put: Callable[[str, str], str]) -> None:
    """An address is checked piece by piece, so a partly replaced address still fails."""
    put("address", p.region)
    put("address", p.road)
    if p.unit:
        put("address", p.unit)


# --- shared resume content ------------------------------------------------------------------


@dataclass
class Layout:
    """Per-file choices for the 인적사항 table."""

    phone: str = "dash"
    birth: str = "dots"
    english: str = "plain"  # plain | upper | reversed
    split_name: bool = False  # 성 / 이름 in separate cells
    split_address: bool = False  # 시·군·구 and 도로명 in separate cells
    address_break: bool = False  # line break inside the address cell
    hanja: bool = False


def english_as(p: Person, style: str) -> str:
    first, last = p.english.rsplit(" ", 1)
    return {
        "plain": p.english,
        "upper": f"{last.upper()} {first.upper().replace('-', '')}",
        "reversed": f"{last}, {first}",
    }[style]


def profile_rows(p: Person, put: Callable[[str, str], str], layout: Layout) -> list[list[str]]:
    """인적사항 rows as cell strings; "\\n" inside a cell is a line break."""
    rows: list[list[str]] = []
    if layout.split_name:
        rows.append(["성", p.surname, "이름", put("candidate", p.given)])
    else:
        rows.append(["성명", put("candidate", p.name), "", ""])
    english = english_as(p, layout.english)
    if layout.english == "upper":
        first = p.english.rsplit(" ", 1)[0].upper().replace("-", "")
        put("candidate", first)
        put("candidate", english)
    else:
        put("candidate", english)
    rows.append(["영문", english, "한자" if layout.hanja else "", put("candidate", p.hanja) if layout.hanja else ""])
    rows.append(["생년월일", put("birthdate", birth_as(p, layout.birth)), "휴대전화", put("phone", phone_as(p, layout.phone))])
    rows.append(["이메일", put("email", p.email), "", ""])
    address_pieces(p, put)
    tail = p.road + (f", {p.unit}" if p.unit else "")
    if layout.split_address:
        rows.append(["주소", p.region, tail, ""])
    elif layout.address_break:
        rows.append(["주소", f"{p.region}\n{tail}", "", ""])
    else:
        rows.append(["주소", p.address, "", ""])
    return rows


def education_rows(p: Person, put: Callable[[str, str], str]) -> list[list[str]]:
    start = int(p.period[:4])
    return [
        ["기간", "학교", "전공", "구분"],
        [f"{start - 4}.03 ~ {start}.02", put("school", p.school), p.major, "졸업"],
    ]


def career_rows(p: Person, put: Callable[[str, str], str]) -> list[list[str]]:
    return [
        ["기간", "회사", "직무", "담당 업무"],
        [p.period, put("company", p.company), p.role, p.job],
    ]


# --- images ---------------------------------------------------------------------------------


def photo_png() -> bytes:
    """ID-photo placeholder: a grey silhouette, no face, no text."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (90, 120), (222, 226, 232))
    draw = ImageDraw.Draw(img)
    draw.ellipse((27, 18, 63, 58), fill=(150, 160, 172))
    draw.ellipse((10, 64, 80, 140), fill=(150, 160, 172))
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


# --- OOXML post-processing ------------------------------------------------------------------


def repack(data: bytes, *, company: str = "", custom: dict[str, str] | None = None) -> bytes:
    """Fixed zip timestamps (stable diffs); optional app.xml Company and custom properties."""
    src = zipfile.ZipFile(io.BytesIO(data))
    files = {name: src.read(name) for name in src.namelist()}
    if "word/comments.xml" in files:  # review comments carry their creation time
        stamp = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode("ascii")
        files["word/comments.xml"] = re.sub(rb'w:date="[^"]*"', b'w:date="' + stamp + b'"', files["word/comments.xml"])
    if company:
        app = files["docProps/app.xml"].decode("utf-8")
        if "<Company>" in app:
            start, end = app.index("<Company>"), app.index("</Company>")
            app = app[: start + len("<Company>")] + company + app[end:]
        elif "<Company/>" in app:
            app = app.replace("<Company/>", f"<Company>{company}</Company>")
        else:
            app = app.replace("</Properties>", f"<Company>{company}</Company></Properties>")
        files["docProps/app.xml"] = app.encode("utf-8")
    if custom:
        props = "".join(
            f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{i}" name="{k}">'
            f"<vt:lpwstr>{v}</vt:lpwstr></property>"
            for i, (k, v) in enumerate(custom.items(), start=2)
        )
        files["docProps/custom.xml"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
            'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
            f"{props}</Properties>"
        ).encode("utf-8")
        types = files["[Content_Types].xml"].decode("utf-8")
        types = types.replace(
            "</Types>",
            '<Override PartName="/docProps/custom.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>',
        )
        files["[Content_Types].xml"] = types.encode("utf-8")
        rels = files["_rels/.rels"].decode("utf-8")
        rels = rels.replace(
            "</Relationships>",
            '<Relationship Id="rIdCustom" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" '
            'Target="docProps/custom.xml"/></Relationships>',
        )
        files["_rels/.rels"] = rels.encode("utf-8")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, content in files.items():
            info = zipfile.ZipInfo(name, date_time=FIXED_TIME.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, content)
    return out.getvalue()


# --- PDF (reportlab) ------------------------------------------------------------------------

FONT = "KR"
_FONT_PATH: list[Path] = []


def register_font(font: Path) -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(FONT, str(font)))
    _FONT_PATH[:] = [font]


def _esc(text: str) -> str:
    from xml.sax.saxutils import escape

    return escape(text).replace("\n", "<br/>")


@dataclass
class PdfPlan:
    layout: Layout = field(default_factory=Layout)
    author: str = ""
    title: str = ""
    subject: str = ""
    keywords: str = ""
    header: str = ""  # drawn at the top of every page
    header_spacing: float = 0  # character spacing (자간); read naively, the text is "권 하 린"
    footer: str = ""  # drawn at the bottom of every page ("{page}" = page number)
    footer_spacing: float = 0
    hidden: tuple[str, ...] = ()  # white text on the white first page
    tiny: tuple[str, ...] = ()  # 0.5 pt text on the first page
    sidebar: tuple[str, ...] = ()  # white lines on a dark sidebar (visible)
    annotation: str = ""  # sticky-note comment on the first page
    email_link: bool = False  # e-mail cell is a mailto: link
    link_label: str = ""  # mailto link with a label instead of the address
    urls: tuple[tuple[str, str], ...] = ()  # (visible text, href)
    narrow_cells: bool = False  # character-wrapped, narrow value column
    vertical_name: bool = False  # one syllable per line in the name cell
    extra_pages: int = 0  # 경력기술서 pages before the 자기소개서


def build_pdf(p: Person, truth: Truth, plan: PdfPlan) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    meta = truth.at("metadata")
    table = truth.at("table")
    link = truth.at("link")

    base = ParagraphStyle("base", fontName=FONT, fontSize=9.5, leading=13.5)
    cell = ParagraphStyle("cell", parent=base, fontSize=9, leading=12)
    narrow = ParagraphStyle("narrow", parent=cell, wordWrap="CJK")
    heading = ParagraphStyle("heading", parent=base, fontSize=13, leading=18, spaceBefore=8, spaceAfter=4)
    title = ParagraphStyle("title", parent=base, fontSize=20, leading=26, alignment=1, spaceAfter=10)

    left = 62 * mm if plan.sidebar else 18 * mm
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=left, rightMargin=18 * mm, topMargin=22 * mm, bottomMargin=20 * mm,
        author=plan.author, title=plan.title, subject=plan.subject, keywords=plan.keywords,
        creator="Hangul Word Processor", invariant=1,
    )  # fmt: skip
    for value, category in ((plan.author, "candidate"), (plan.title, None), (plan.subject, None), (plan.keywords, None)):
        if value and category:
            meta(category, value)

    def say(canvas, x, y, line, size, spacing=0.0, centred=False) -> None:
        """One line through a text object, so character spacing never leaks to the next."""
        if centred:
            x -= (canvas.stringWidth(line, FONT, size) + spacing * len(line)) / 2
        text = canvas.beginText(x, y)
        text.setFont(FONT, size)
        text.setCharSpace(spacing)
        text.textOut(line)
        canvas.drawText(text)

    def page_decor(canvas, document) -> None:
        canvas.saveState()
        width, height = A4
        if plan.sidebar:
            canvas.setFillColorRGB(0.13, 0.18, 0.30)
            canvas.rect(0, 0, 54 * mm, height, fill=1, stroke=0)
        canvas.setFillColorRGB(0.35, 0.35, 0.35)
        if plan.header:
            say(canvas, left, height - 12 * mm, plan.header, 8, plan.header_spacing)
        if plan.footer:
            centre = width / 2 + (left - 18 * mm) / 2
            say(canvas, centre, 10 * mm, plan.footer.format(page=document.page), 8, plan.footer_spacing, centred=True)
        if document.page == 1:
            canvas.setFillColorRGB(1, 1, 1)
            y = height - 40 * mm
            for i, line in enumerate(plan.sidebar):
                say(canvas, 8 * mm, y, line, 13 if i == 0 else 8.5)
                y -= 9 * mm if i == 0 else 6 * mm
            for i, line in enumerate(plan.hidden):
                say(canvas, left, 15 * mm - i * 3 * mm, line, 7)
            canvas.setFillColorRGB(0, 0, 0)
            for i, line in enumerate(plan.tiny):
                say(canvas, width - 40 * mm, 16 * mm - i * 2 * mm, line, 0.5)
            if plan.annotation:
                rect = (width - 30 * mm, height - 30 * mm, width - 22 * mm, height - 22 * mm)
                canvas.textAnnotation(plan.annotation, Rect=rect, relative=0)
        canvas.restoreState()

    story: list = [Paragraph("이 력 서", title)]

    # 인적사항: photo column spanning every row, then label/value pairs.
    rows = profile_rows(p, table, plan.layout)
    if plan.email_link:
        link("email", p.email)
    data, spans = [], []
    for r, row in enumerate(rows):
        cells = []
        for c, value in enumerate(row):
            style = narrow if plan.narrow_cells and c in (1, 3) else cell
            if plan.vertical_name and r == 0 and c == 1:
                value = "\n".join(value)
            if plan.email_link and value == p.email:
                cells.append(Paragraph(f'<a href="mailto:{p.email}" color="blue">{_esc(value)}</a>', style))
            else:
                cells.append(Paragraph(_esc(value), style))
        data.append(["" if r else Image(io.BytesIO(photo_png()), width=24 * mm, height=32 * mm), *cells])
        if row[2] == "" and row[3] == "":
            spans.append(("SPAN", (2, r), (4, r)))
    truth.images += 1
    widths = [28 * mm, 18 * mm, 22 * mm, 16 * mm, 22 * mm] if plan.narrow_cells else [28 * mm, 18 * mm, 48 * mm, 18 * mm, 40 * mm]
    if plan.sidebar:
        widths = [26 * mm, 16 * mm, 40 * mm, 16 * mm, 32 * mm]
    profile = Table(data, colWidths=widths)
    profile.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("SPAN", (0, 0), (0, len(data) - 1)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (1, 0), (1, -1), colors.whitesmoke),
        *spans,
    ]))  # fmt: skip
    story += [profile, Spacer(1, 6 * mm)]

    for heading_text, rows in (("학력 사항", education_rows(p, table)), ("경력 사항", career_rows(p, table))):
        story.append(Paragraph(heading_text, heading))
        grid = Table([[Paragraph(_esc(v), narrow if plan.narrow_cells else cell) for v in row] for row in rows],
                     colWidths=[34 * mm, 40 * mm, 34 * mm, 38 * mm] if not plan.sidebar else [30 * mm, 34 * mm, 30 * mm, 36 * mm])  # fmt: skip
        grid.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke)]))
        story += [grid, Spacer(1, 4 * mm)]

    if plan.urls or plan.link_label:
        parts = []
        for text, href in plan.urls:
            parts.append(f'<a href="{href}" color="blue">{_esc(text)}</a>')
        if plan.link_label:
            link("email", p.email)
            parts.append(f'<a href="mailto:{p.email}" color="blue">{_esc(plan.link_label)}</a>')
        story.append(Paragraph("링크: " + " · ".join(parts), base))

    for n in range(plan.extra_pages):
        story += [PageBreak(), Paragraph(f"경력기술서 ({n + 1})", heading)]
        for k in range(6):
            story.append(Paragraph(
                f"프로젝트 {n * 6 + k + 1}. 요구사항을 정리하고 설계 문서를 작성한 뒤 2주 단위로 배포했습니다. "
                "성능 지표와 장애 건수를 매주 공유했고, 회고에서 나온 개선 항목을 다음 주기에 반영했습니다.", base))  # fmt: skip
            story.append(Spacer(1, 3 * mm))

    story += [PageBreak(), Paragraph("자기소개서", heading)]
    for paragraph in intro(p, truth, "body"):
        story += [Paragraph(_esc(paragraph), base), Spacer(1, 3 * mm)]

    doc.build(story, onFirstPage=page_decor, onLaterPages=page_decor)
    if plan.header:
        truth.trap("머리글·바닥글 반복")
    return buf.getvalue()


def pdf_r01(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("파일명에 이름", "작성자 메타데이터에 실명", "증명사진", "자간 넓힌 머리글 이름·바닥글 전화번호(권 하 린, 0 1 0 . 0 2 ...)", "표 셀 안 줄바꿈 주소",
               "영문·한자 이름", "mailto 링크", "전화번호 여러 형식", "흰 글씨 숨은 텍스트", "GitHub 주소에 영문 이름")  # fmt: skip
    hidden_text = f"키워드 {p.name} {phone_as(p, 'plain')} 백엔드 우대"
    truth.put("hidden", "candidate", p.name)
    truth.put("hidden", "phone", phone_as(p, "plain"))
    truth.put("body", "candidate", p.name)  # header (spaced by character spacing)
    truth.put("body", "phone", phone_as(p, "dots"))  # footer
    truth.put("body", "candidate", "harin-kwon")
    truth.put("link", "candidate", "harin-kwon")
    plan = PdfPlan(
        layout=Layout(phone="dash", birth="korean", address_break=True, hanja=True),
        author=p.name, title=f"이력서 - {p.name}", header=f"{p.name} · 이력서", header_spacing=5,
        footer=f"{p.name} | {phone_as(p, 'dots')} | {{page}}", footer_spacing=4, hidden=(hidden_text,), email_link=True,
        urls=(("github.com/harin-kwon", "https://github.com/harin-kwon"),),
    )  # fmt: skip
    truth.put("metadata", "candidate", p.name)
    return build_pdf(p, truth, plan), {}


def pdf_r02(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("복성(남궁)", "파일명에 이름", "제목·키워드 메타데이터", "띄어 쓴 이름(남궁 민재)", "영문 대문자 이름", "성/이름 칸 분리",
               "+82 국제 형식 전화번호", "PDF 주석(메모)에 연락처", "증명사진")  # fmt: skip
    truth.put("metadata", "candidate", f"{p.surname} {p.given}")
    truth.put("metadata", "email", p.email)
    truth.put("body", "candidate", f"{p.surname} {p.given}")
    note = f"검토 메모: {p.name} 지원자 연락처 {phone_as(p, 'dash')} 재확인 필요"
    truth.put("comment", "candidate", p.name)
    truth.put("comment", "phone", phone_as(p, "dash"))
    plan = PdfPlan(
        layout=Layout(phone="intl", birth="dots", english="upper", split_name=True),
        title=f"{p.surname} {p.given} 입사지원서", keywords=f"{p.surname} {p.given}, 데이터 분석, {p.email}",
        header=f"입사지원서 · {p.surname} {p.given}", footer="- {page} -", annotation=note,
    )  # fmt: skip
    return build_pdf(p, truth, plan), {}


def pdf_r03(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", "SerinOh")
    truth.trap("영문 파일명(SerinOh)", "어두운 사이드바의 흰 글씨(보이는 텍스트)", "2단 레이아웃", "점 구분 전화번호", "mailto 링크(문구만 보임)",
               "GitHub 주소에 영문 이름", "증명사진")  # fmt: skip
    body = truth.at("body")
    sidebar = (
        body("candidate", p.name),
        body("candidate", p.english),
        body("phone", phone_as(p, "dots")),
        body("email", p.email),
        p.region,
        p.road,
    )
    address_pieces(p, body)
    truth.put("body", "candidate", "serin-oh")
    truth.put("link", "candidate", "serin-oh")
    plan = PdfPlan(
        layout=Layout(phone="dots", birth="slash"), author="Serin Oh", title="Resume",
        sidebar=sidebar, link_label="메일 보내기", urls=(("github.com/serin-oh", "https://github.com/serin-oh"),),
    )  # fmt: skip
    truth.put("metadata", "candidate", "Serin Oh")
    return build_pdf(p, truth, plan), {}


def pdf_r04(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("4쪽 문서에 바닥글 반복(이름·전화)", "영문 작성자 메타데이터", "두 자리 연도 생년월일", "학교 약칭(카이스트)", "증명사진",
               "자간 넓힌 머리글에 지원서에 없는 추천인 이름·전화(추 천 인 정 민 지 박 사 0 1 0 - ...)")  # fmt: skip
    truth.put("metadata", "candidate", p.english)
    truth.put("body", "candidate", p.name)
    truth.put("body", "phone", phone_as(p, "spaces"))
    ref_phone = truth.put("body", "phone", "010-0374-2281")
    truth.put("body", "person", p.ref)
    truth.keep.append("추천인")
    plan = PdfPlan(
        layout=Layout(phone="paren", birth="short"), author=p.english, title="자기소개서 최종",
        header=f"추천인 {p.ref} {p.ref_title} {ref_phone}", header_spacing=4,
        footer=f"{p.name} / {phone_as(p, 'spaces')} / {{page}}쪽", extra_pages=2,
    )  # fmt: skip
    return build_pdf(p, truth, plan), {}


def pdf_r05(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    upload = unicodedata.normalize("NFD", p.file)
    truth.put("filename", "candidate", p.name)
    truth.trap("macOS 업로드 파일명(NFD 분해 한글)", "복성(황보)", "전각 숫자 전화번호", "0.5pt 숨은 텍스트", "흰 글씨 숨은 텍스트",
               "호칭 붙은 추천인 이름", "증명사진")  # fmt: skip
    tiny = f"{p.name} {p.email}"
    truth.put("hidden", "candidate", p.name)
    truth.put("hidden", "email", p.email)
    hidden = f"추천인 {p.ref} {p.ref_title} 확인 완료"
    truth.put("hidden", "person", p.ref)
    plan = PdfPlan(
        layout=Layout(phone="fullwidth", birth="dots"), author=p.name, tiny=(tiny,), hidden=(hidden,),
        header=f"{p.name} 지원서",
    )  # fmt: skip
    truth.put("metadata", "candidate", p.name)
    truth.put("body", "candidate", p.name)
    return build_pdf(p, truth, plan), {"upload_name": upload}


def pdf_r06(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("좁은 표 셀에서 글자 단위 줄바꿈(주소·학교·회사)", "세로로 쪼갠 이름 칸", "키워드 메타데이터에 이메일", "공백 구분 전화번호",
               "(주) 붙은 회사명", "증명사진")  # fmt: skip
    truth.put("metadata", "email", p.email)
    plan = PdfPlan(
        layout=Layout(phone="spaces", birth="iso"), keywords=f"임베디드, {p.email}", title="입사지원서",
        narrow_cells=True, vertical_name=True, footer="{page}",
    )  # fmt: skip
    return build_pdf(p, truth, plan), {}


# --- scanned PDFs (text rendered into images, no text layer) ----------------------------------


def _scan_image(lines: list[str], seed: int) -> bytes:
    import random

    from PIL import Image, ImageDraw, ImageFont

    width, height = 827, 1169  # A4 at 100 dpi
    img = Image.new("L", (width, height), 250)
    draw = ImageDraw.Draw(img)
    regular = ImageFont.truetype(str(_FONT_PATH[0]), 17)
    large = ImageFont.truetype(str(_FONT_PATH[0]), 30)
    y = 70
    for line in lines:
        if line.startswith("# "):
            draw.text((width // 2, y), line[2:], fill=25, font=large, anchor="mt")
            y += 56
        elif line == "---":
            draw.line((70, y, width - 70, y), fill=90, width=1)
            y += 14
        else:
            draw.text((80, y), line, fill=30, font=regular)
            y += 27
    rng = random.Random(seed)
    for _ in range(1500):  # scanner speckle
        x, yy = rng.randrange(width), rng.randrange(height)
        draw.point((x, yy), fill=rng.randrange(120, 230))
    img = img.rotate(0.7, fillcolor=250, resample=Image.Resampling.BICUBIC)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=55, optimize=True)
    return buf.getvalue()


def _wrap(text: str, width: int = 44) -> list[str]:
    out, line = [], ""
    for word in text.split():
        if len(line) + len(word) + 1 > width:
            out.append(line)
            line = word
        else:
            line = f"{line} {word}".strip()
    return out + ([line] if line else [])


def _scan_lines(p: Person, truth: Truth, with_intro: bool) -> list[str]:
    img = truth.at("image")
    lines = [
        "# 이 력 서",
        f"성명: {img('candidate', p.name)}  ({img('candidate', p.english)})",
        f"생년월일: {img('birthdate', birth_as(p, 'dots'))}",
        f"휴대전화: {img('phone', p.phone)}    이메일: {img('email', p.email)}",
        f"주소: {p.address}",
        "---",
        f"학력: {img('school', p.school)} {p.major} 졸업",
        f"경력: {img('company', p.company)} {p.role} ({p.period})",
        "---",
    ]
    address_pieces(p, img)
    if with_intro:
        lines.append("자기소개서")
        for paragraph in intro(p, truth, "image"):
            lines += _wrap(paragraph) + [""]
    return lines


def _scan_pdf(p: Person, truth: Truth, pages: list[bytes | list[str]], *, author: str = "", watermark: str = "") -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    if author:
        c.setAuthor(author)
    c.setCreator("ScanApp 3.2")
    for page in pages:
        if isinstance(page, bytes):
            c.drawImage(ImageReader(io.BytesIO(page)), 0, 0, width=A4[0], height=A4[1])
            truth.images += 1
            if watermark:
                c.setFont(FONT, 6)
                c.setFillColorRGB(0.5, 0.5, 0.5)
                c.drawString(20, 12, watermark)
        else:  # a typed page with a real text layer
            text = c.beginText(60, A4[1] - 70)
            text.setFont(FONT, 10)
            text.setLeading(15)
            for line in page:
                text.textLine(line)
            c.drawText(text)
        c.showPage()
    c.save()
    return buf.getvalue()


def scan_r18(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.put("metadata", "candidate", p.name)
    truth.trap("텍스트 레이어 없는 1쪽 스캔", "작성자 메타데이터에 실명(보류라 전송 안 됨)", "파일명에 이름")
    page = _scan_image(_scan_lines(p, truth, with_intro=True), seed=18)
    return _scan_pdf(p, truth, [page], author=p.name), {"ocr_pages": [1]}


def scan_r19(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("2쪽 스캔", "스캔 앱 워터마크 글자만 있는 텍스트 레이어")
    lines = _scan_lines(p, truth, with_intro=True)
    first, second = lines[:9], lines[9:]
    pages = [_scan_image(first, seed=191), _scan_image(["# 자기소개서", *second], seed=192)]
    return _scan_pdf(p, truth, pages, watermark="Scanned with ScanApp"), {"ocr_pages": [1, 2]}


def scan_r20(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.put("metadata", "candidate", p.name)
    truth.trap("혼합 문서: 1쪽은 텍스트, 2쪽만 스캔", "일부만 읽히면 전체 보류")
    typed = ["자기소개서", ""]
    for paragraph in intro(p, truth, "body"):
        typed += _wrap(paragraph, 48) + [""]
    scanned = _scan_image(_scan_lines(p, truth, with_intro=False), seed=20)
    return _scan_pdf(p, truth, [typed, scanned], author=p.name), {"ocr_pages": [2]}


# --- DOCX (python-docx) ---------------------------------------------------------------------


@dataclass
class DocxPlan:
    layout: Layout = field(default_factory=Layout)
    author: str = ""
    last_modified_by: str = ""
    title: str = ""
    comments: str = ""  # core property "comments" (설명)
    company: str = ""  # app.xml
    custom: dict[str, str] = field(default_factory=dict)  # custom.xml
    header: str = ""
    first_header: str = ""  # different first page
    footer: str = ""
    email_link: bool = False
    white: str = ""  # white run after the 자기소개서
    vanish: str = ""  # hidden (w:vanish) run
    comment: str = ""  # review comment anchored on the name cell
    nested: bool = False  # 연락처 as a table nested in a cell
    page_breaks: int = 0
    extra: tuple[str, ...] = ()  # paragraphs after the 자기소개서


def _hyperlink(paragraph, url: str, text: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def build_docx(p: Person, truth: Truth, plan: DocxPlan) -> bytes:
    import docx
    from docx.shared import Cm, Pt, RGBColor

    table_put = truth.at("table")
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10)

    props = doc.core_properties
    props.author, props.last_modified_by, props.title = plan.author, plan.last_modified_by, plan.title
    props.comments, props.created, props.modified, props.revision = plan.comments, FIXED_TIME, FIXED_TIME, 3

    section = doc.sections[0]
    if plan.header:
        section.header.paragraphs[0].text = plan.header
    if plan.first_header:
        section.different_first_page_header_footer = True
        section.first_page_header.paragraphs[0].text = plan.first_header
    if plan.footer:
        section.footer.paragraphs[0].text = plan.footer

    doc.add_heading("이 력 서", level=0)
    rows = profile_rows(p, table_put, plan.layout)
    grid = doc.add_table(rows=len(rows), cols=5)
    grid.style = "Table Grid"
    photo = grid.cell(0, 0).merge(grid.cell(len(rows) - 1, 0))
    photo.paragraphs[0].add_run().add_picture(io.BytesIO(photo_png()), width=Cm(2.4))
    truth.images += 1
    name_run = None
    for r, row in enumerate(rows):
        for c, value in enumerate(row, start=1):
            cell = grid.cell(r, c)
            if not value:
                continue
            paragraph = cell.paragraphs[0]
            if plan.email_link and value == p.email:
                _hyperlink(paragraph, f"mailto:{p.email}", value)
                truth.put("link", "email", p.email)
            elif plan.nested and value == p.email:
                inner = cell.add_table(rows=2, cols=2)
                inner.style = "Table Grid"
                inner.cell(0, 0).text, inner.cell(0, 1).text = "메일", value
                inner.cell(1, 0).text, inner.cell(1, 1).text = "블로그", "haram-sunwoo.example.net"
                truth.put("table", "candidate", "haram-sunwoo")
            else:
                for i, line in enumerate(value.split("\n")):
                    run = paragraph.add_run(line)
                    if i < len(value.split("\n")) - 1:
                        run.add_break()
                    if r == 0 and c == 2:
                        name_run = run
    if plan.comment and name_run is not None:
        doc.add_comment(name_run, text=plan.comment, author="인사팀", initials="HR")

    for heading, rows in (("학력 사항", education_rows(p, table_put)), ("경력 사항", career_rows(p, table_put))):
        doc.add_heading(heading, level=2)
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                t.cell(r, c).text = value

    for n in range(plan.page_breaks):
        doc.add_page_break()
        doc.add_heading(f"경력기술서 ({n + 1})", level=2)
        for k in range(5):
            doc.add_paragraph(
                f"프로젝트 {n * 5 + k + 1}. 요구사항 분석부터 운영까지 맡았고, 주간 지표를 공유하며 개선 항목을 다음 배포에 반영했습니다."
            )

    doc.add_page_break()
    doc.add_heading("자기소개서", level=1)
    for paragraph in intro(p, truth, "body"):
        doc.add_paragraph(paragraph)
    for paragraph in plan.extra:
        doc.add_paragraph(paragraph)
    if plan.white or plan.vanish:
        paragraph = doc.add_paragraph("이상입니다.")
        if plan.white:
            run = paragraph.add_run(" " + plan.white)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(4)
        if plan.vanish:
            run = paragraph.add_run(" " + plan.vanish)
            run.font.hidden = True

    buf = io.BytesIO()
    doc.save(buf)
    return repack(buf.getvalue(), company=plan.company, custom=plan.custom)


def docx_r07(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("머리글(이름·전화)·바닥글(이름)", "작성자·마지막 수정자 메타데이터", "app.xml 회사명", "증명사진", "셀 안 줄바꿈 주소",
               "mailto 하이퍼링크", "흰 글씨 숨은 텍스트")  # fmt: skip
    header = f"{p.name} | 지원 직무: {p.job} | {phone_as(p, 'dash')}"
    truth.put("header", "candidate", p.name)
    truth.put("header", "phone", phone_as(p, "dash"))
    truth.put("footer", "candidate", p.name)
    truth.put("metadata", "candidate", p.name)
    truth.put("metadata", "candidate", p.english)
    truth.put("metadata", "company", p.company)
    white = f"{p.name} {phone_as(p, 'plain')} {p.email}"
    truth.put("hidden", "candidate", p.name)
    truth.put("hidden", "phone", phone_as(p, "plain"))
    truth.put("hidden", "email", p.email)
    plan = DocxPlan(
        layout=Layout(phone="dash", birth="korean", address_break=True), author=p.name, last_modified_by=p.english,
        title="이력서", company=p.company, header=header, footer=f"{p.name} 이력서", email_link=True, white=white,
    )  # fmt: skip
    return build_docx(p, truth, plan), {}


def docx_r08(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("첫 페이지만 다른 머리글", "숨김 서식(w:vanish) 텍스트", "검토 메모(comment)", "주소를 두 칸으로 분리", "+82 (0) 전화번호",
               "설명 메타데이터에 이메일", "증명사진")  # fmt: skip
    truth.put("header", "candidate", p.name)
    truth.put("header", "email", p.email)
    truth.put("header", "phone", phone_as(p, "dots"))
    truth.put("metadata", "email", p.email)
    truth.put("comment", "candidate", p.given)
    truth.put("comment", "phone", phone_as(p, "plain"))
    vanish = f"예전 주소 {p.region} {p.road}"
    address_pieces(p, truth.at("hidden"))
    plan = DocxPlan(
        layout=Layout(phone="intl0", birth="slash", split_address=True), author="HR-PC", title="입사지원서",
        comments=f"문의: {p.email}", first_header=f"입사지원서 · {p.name} · {p.email}", header=f"{p.name} · {phone_as(p, 'dots')}",
        comment=f"{p.given} 님 경력 기간 확인 필요, 연락처 {phone_as(p, 'plain')}", vanish=vanish, page_breaks=1,
    )  # fmt: skip
    return build_docx(p, truth, plan), {}


def docx_r09(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("복성(선우)", "성/이름 칸 분리", "하이픈 영문 이름(Ha-ram Sunwoo)", "표 안의 표(중첩)", "블로그 주소에 영문 이름",
               "작성자 메타데이터 영문", "증명사진", "지원서에 없는 추천인 전화·이메일을 글자마다 띄어 씀(0 1 0 - 1 8 ..., s j . y o o n @ ...)")  # fmt: skip
    truth.put("metadata", "candidate", "Haram Sunwoo")
    truth.put("footer", "candidate", "Ha-ram Sunwoo")
    truth.put("footer", "email", p.email)
    ref_phone = truth.put("body", "phone", "010-1826-4097")
    ref_email = truth.put("body", "email", "sj.yoon@example.net")
    truth.keep.append("추천인 연락처:")
    plan = DocxPlan(
        layout=Layout(phone="plain", birth="dots", split_name=True), author="Haram Sunwoo", title="Resume",
        footer=f"Ha-ram Sunwoo · {p.email}", nested=True,
        extra=(f"추천인 연락처: {' '.join(ref_phone)} / {' '.join(ref_email)}",),
    )  # fmt: skip
    truth.put("table", "email", p.email)
    return build_docx(p, truth, plan), {}


def docx_r10(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("사용자 정의 문서 속성(custom.xml)", "4쪽 문서 바닥글 반복", "한자 이름", "영문 이름 성·이름 순서 뒤집힘", "기관 약칭(유니스트·현대차)",
               "증명사진")  # fmt: skip
    truth.put("metadata", "candidate", p.name)
    truth.put("metadata", "phone", p.phone)
    truth.put("footer", "candidate", p.name)
    truth.put("footer", "phone", phone_as(p, "intl"))
    plan = DocxPlan(
        layout=Layout(phone="spaces", birth="iso", english="reversed", hanja=True), author=p.name, title="이력서",
        custom={"지원자": p.name, "연락처": p.phone}, footer=f"{p.name} · {phone_as(p, 'intl')}", page_breaks=2,
    )  # fmt: skip
    return build_docx(p, truth, plan), {}


# --- PPTX (python-pptx) ---------------------------------------------------------------------


def _textbox(shapes, x, y, w, h, lines, size=14, color=None, bold_first=False):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    box = shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size + (6 if bold_first and i == 0 else 0))
        if color is not None:
            run.font.color.rgb = RGBColor(*color)
    return box


def _pptx_table(slide, x, y, w, rows, put=None):
    from pptx.util import Inches, Pt

    shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(0.4) * len(rows))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = shape.table.cell(r, c)
            cell.text = value.replace("\n", "\v")
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(12)
    return shape


def build_pptx(p: Person, truth: Truth, plan: dict) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    props = prs.core_properties
    props.author = plan.get("author", "")
    props.last_modified_by = plan.get("last_modified_by", "")
    props.title = plan.get("title", "")
    props.created = props.modified = FIXED_TIME
    props.revision = 2
    blank = prs.slide_layouts[6]

    # 1. cover
    cover = prs.slides.add_slide(blank)
    title_lines = plan["cover"]
    _textbox(cover.shapes, Inches(1), Inches(2.3), Inches(11), Inches(2.5), title_lines, size=24, bold_first=True)

    # 2. 인적사항
    profile = prs.slides.add_slide(blank)
    _textbox(profile.shapes, Inches(0.6), Inches(0.4), Inches(8), Inches(0.8), ["인적사항"], size=26)
    profile.shapes.add_picture(io.BytesIO(photo_png()), Inches(0.6), Inches(1.4), Inches(1.8), Inches(2.4))
    truth.images += 1
    rows = profile_rows(p, truth.at("table"), plan.get("layout", Layout()))
    _pptx_table(profile, Inches(2.8), Inches(1.4), Inches(9.8), rows)

    # 3. 학력·경력
    history = prs.slides.add_slide(blank)
    _textbox(history.shapes, Inches(0.6), Inches(0.4), Inches(8), Inches(0.8), ["학력 · 경력"], size=26)
    _pptx_table(history, Inches(0.6), Inches(1.4), Inches(12), education_rows(p, truth.at("table")))
    _pptx_table(history, Inches(0.6), Inches(3.2), Inches(12), career_rows(p, truth.at("table")))

    # 4. 자기소개서
    letter = prs.slides.add_slide(blank)
    _textbox(letter.shapes, Inches(0.6), Inches(0.4), Inches(8), Inches(0.8), ["자기소개서"], size=26)
    _textbox(letter.shapes, Inches(0.6), Inches(1.3), Inches(12), Inches(5.5), intro(p, truth, "body"), size=13)

    slides = [cover, profile, history, letter]
    for hook in plan.get("hooks", ()):
        hook(prs, slides)

    buf = io.BytesIO()
    prs.save(buf)
    return repack(buf.getvalue())


def pptx_r11(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    from pptx.util import Inches

    truth.put("filename", "candidate", p.name)
    truth.trap("슬라이드 노트에 이름·전화", "슬라이드 마스터 바닥글에 이름(모든 슬라이드 반복)", "증명사진", "텍스트 링크(mailto)", "작성자 메타데이터",
               "흰 글씨 숨은 텍스트")  # fmt: skip
    truth.put("metadata", "candidate", p.name)
    for value, category in ((p.name, "candidate"), (p.english, "candidate")):
        truth.put("body", category, value)
    truth.put("footer", "candidate", p.name)
    truth.put("footer", "email", p.email)
    truth.put("notes", "candidate", p.given)
    truth.put("notes", "phone", phone_as(p, "dash"))
    truth.put("link", "email", p.email)
    truth.put("hidden", "candidate", p.name)
    truth.put("hidden", "phone", phone_as(p, "plain"))

    def hooks(prs, slides):
        cover, profile, history, letter = slides
        # Master footer: a text box on the slide master shows on every slide.
        tmp = _textbox(cover.shapes, Inches(0.4), Inches(7.0), Inches(8), Inches(0.4), [f"{p.name} 포트폴리오 · {p.email}"], size=10)
        tmp._element.xpath("./p:nvSpPr/p:cNvPr")[0].set("id", "900")
        prs.slide_master.shapes._spTree.append(tmp._element)
        # mailto link on a run
        box = _textbox(letter.shapes, Inches(0.6), Inches(6.6), Inches(6), Inches(0.5), ["연락하기"], size=12)
        box.text_frame.paragraphs[0].runs[0].hyperlink.address = f"mailto:{p.email}"
        # white text on a box without fill
        _textbox(history.shapes, Inches(9), Inches(6.5), Inches(4), Inches(0.5), [f"{p.name} {phone_as(p, 'plain')}"], size=8, color=(255, 255, 255))
        for slide, note in ((cover, f"{p.given} 님 발표 시작: 인사 후 경력 요약"), (profile, f"질문 받으면 {phone_as(p, 'dash')}로 연락 드린다고 안내")):
            slide.notes_slide.notes_text_frame.text = note

    plan = {"author": p.name, "title": "포트폴리오", "cover": ["포트폴리오", f"{p.name} ({p.english})", p.job], "hooks": [hooks],
            "layout": Layout(phone="dash", birth="dots")}  # fmt: skip
    return build_pptx(p, truth, plan), {}


def pptx_r12(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    from pptx.util import Inches

    truth.put("filename", "candidate", "MinhyukPyo")
    truth.trap("영문 파일명", "그룹 도형 안 주소", "슬라이드 밖에 놓인 텍스트 상자(전화)", "도형 클릭 하이퍼링크(mailto)", "마지막 수정자 메타데이터",
               "노트에 추천인", "증명사진")  # fmt: skip
    truth.put("metadata", "candidate", p.english)
    truth.put("body", "candidate", p.name)
    truth.put("body", "candidate", p.english)
    address_pieces(p, truth.at("body"))
    truth.put("hidden", "phone", phone_as(p, "paren"))
    truth.put("hidden", "candidate", p.name)
    truth.put("link", "email", p.email)
    truth.put("notes", "person", p.ref)

    def hooks(prs, slides):
        cover, profile, history, letter = slides
        group = profile.shapes.add_group_shape()
        _textbox(group.shapes, Inches(2.8), Inches(5.0), Inches(4), Inches(0.5), ["거주지"], size=11)
        _textbox(group.shapes, Inches(2.8), Inches(5.4), Inches(6), Inches(0.5), [f"{p.region} {p.road}"], size=11)
        _textbox(cover.shapes, prs.slide_width + Inches(1), Inches(1), Inches(4), Inches(0.6), [f"{p.name} 개인 연락처 {phone_as(p, 'paren')}"], size=12)
        button = _textbox(letter.shapes, Inches(10.5), Inches(6.6), Inches(2.4), Inches(0.5), ["메일 문의"], size=12)
        button.click_action.hyperlink.address = f"mailto:{p.email}"
        letter.notes_slide.notes_text_frame.text = f"추천인: {p.ref} {p.ref_title}님 (전 직장 상사)"

    plan = {"author": "PC", "last_modified_by": p.english, "title": "Portfolio", "cover": ["Portfolio", f"{p.name} · {p.english}"],
            "hooks": [hooks], "layout": Layout(phone="dash", birth="korean")}  # fmt: skip
    return build_pptx(p, truth, plan), {}


def pptx_r13(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("숨김 슬라이드", "성/이름 칸 분리", "제목 메타데이터에 이름", "노트에 추천인과 이메일", "시·군·구 없는 주소(세종)", "증명사진")
    truth.put("metadata", "candidate", p.name)
    truth.put("notes", "person", p.ref)
    truth.put("notes", "email", p.email)
    truth.put("body", "candidate", p.name)
    truth.put("body", "phone", phone_as(p, "dots"))

    def hooks(prs, slides):
        from pptx.util import Inches

        cover, profile, history, letter = slides
        spare = prs.slides.add_slide(prs.slide_layouts[6])
        _textbox(spare.shapes, Inches(0.6), Inches(0.6), Inches(12), Inches(2), ["(발표용 아님) 연락처", f"{p.name} {phone_as(p, 'dots')}"], size=16)
        spare._element.set("show", "0")
        letter.notes_slide.notes_text_frame.text = f"{p.ref} {p.ref_title}님이 추천서를 {p.email}로 보내 주실 예정"

    plan = {"title": f"{p.name} 경력기술서", "cover": ["경력기술서", p.job], "hooks": [hooks],
            "layout": Layout(phone="dash", birth="dots", split_name=True)}  # fmt: skip
    return build_pptx(p, truth, plan), {}


# --- HTML -----------------------------------------------------------------------------------


def _h(text: str) -> str:
    from html import escape

    return escape(text, quote=True).replace("\n", "<br>")


def _photo_data_uri() -> str:
    import base64

    return "data:image/png;base64," + base64.b64encode(photo_png()).decode("ascii")


def _html_rows(rows: list[list[str]], *, email_link: str = "") -> str:
    out = []
    for row in rows:
        cells = []
        for c in range(0, len(row), 2):
            label, value = row[c], row[c + 1]
            if not label and not value:
                continue
            span = ' colspan="3"' if c == 0 and not row[2] and not row[3] else ""
            shown = f'<a href="mailto:{value}">{_h(value)}</a>' if email_link and value == email_link else _h(value)
            cells.append(f"<th>{_h(label)}</th><td{span}>{shown}</td>")
        out.append("<tr>" + "".join(cells) + "</tr>")
    return "\n".join(out)


def _html_grid(rows: list[list[str]]) -> str:
    head = "".join(f"<th>{_h(v)}</th>" for v in rows[0])
    body = "".join("<tr>" + "".join(f"<td>{_h(v)}</td>" for v in row) + "</tr>" for row in rows[1:])
    return f"<table border=\"1\"><tr>{head}</tr>{body}</table>"


def html_r14(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("<title>·<meta author>에 이름", "description 메타에 전화번호", "<header>·<footer> 요소", "mailto·tel 링크",
               "증명사진(data URI, alt에 이름)", "<br>로 쪼갠 주소", "띄어 쓴 미지의 학교명(가 람 대 학 교)",
               "@·점 앞뒤를 띄운 추천인 이메일(doyoon.nam @ example . com)")  # fmt: skip
    meta = truth.at("metadata")
    tel = "+82" + p.phone_digits[1:]
    rows = profile_rows(p, truth.at("table"), Layout(phone="dash", birth="korean", address_break=True))
    truth.put("link", "email", p.email)
    truth.put("link", "phone", tel)
    truth.images += 1
    letter = "\n".join(f"<p>{_h(x)}</p>" for x in intro(p, truth, "body"))
    exchange = " ".join(truth.put("body", "school", "가람대학교"))
    local, domain = truth.put("body", "email", "doyoon.nam@example.com").split("@")
    ref_email = f"{local} @ {domain.replace('.', ' . ')}"
    truth.put("body", "person", p.ref)
    truth.keep.extend(["교환학생", "추천인"])
    html = f"""<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>이력서 - {meta('candidate', p.name)}</title>
<meta name="author" content="{p.name}">
<meta name="description" content="{p.job} 지원자 이력서. 연락처 {meta('phone', p.phone)}">
<style>body {{ font-family: sans-serif; max-width: 760px; margin: 2rem auto; }} th {{ background: #f3f4f6; }}</style>
</head>
<body>
<header><strong>{truth.put('header', 'candidate', p.name)}</strong> · {_h(p.job)} 지원 · <a href="tel:{tel}">전화하기</a></header>
<main>
<h1>이 력 서</h1>
<img src="{_photo_data_uri()}" alt="{p.name} 증명사진" width="90" height="120">
<table border="1" class="profile">
{_html_rows(rows, email_link=p.email)}
</table>
<h2>학력</h2>
{_html_grid(education_rows(p, truth.at("table")))}
<h2>경력</h2>
{_html_grid(career_rows(p, truth.at("table")))}
<h2>자기소개서</h2>
{letter}
<h2>기타</h2>
<p>교환학생: {exchange} (2017)</p>
<p>추천인 {p.ref} {p.ref_title} 메일: {ref_email}</p>
</main>
<footer>&copy; 2026 {truth.put('footer', 'candidate', p.name)} · {truth.put('footer', 'email', p.email)}</footer>
</body>
</html>
"""
    return html.encode("utf-8"), {}


def html_r15(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", "jegal_yunho")
    truth.trap("EUC-KR 인코딩(CP949 확장 글자 포함)", "복성(제갈)", "&nbsp;로 띄운 이름", "CSS 클래스(display:none) 숨은 연락처",
               "영문 파일명(성_이름 순서)", "흰 배경 아닌 칸의 흰 글씨(보이는 텍스트)")  # fmt: skip
    truth.put("metadata", "candidate", p.name)
    truth.put("body", "candidate", f"{p.surname} {p.given}")
    hidden = truth.at("hidden")
    memo = f"비공개 메모: {hidden('candidate', p.name)} {hidden('phone', p.phone)} {hidden('email', p.email)}"
    rows = profile_rows(p, truth.at("table"), Layout(phone="dots", birth="dots"))
    letter = "\n".join(f"<p>{_h(x)}</p>" for x in intro(p, truth, "body"))
    truth.keep.append("똠양꿍")
    html = f"""<html>
<head>
<meta http-equiv="Content-Type" content="text/html; charset=euc-kr">
<title>{p.name} 입사지원서</title>
<style type="text/css">
.memo {{ display: none }}
td {{ font-size: 12px }}
.banner {{ background: #1f2d4d; color: #ffffff }}
</style>
</head>
<body bgcolor="#ffffff">
<div class="banner">입사지원서 - {_h(p.job)}</div>
<p><font size="5">성명: {p.surname}&nbsp;{p.given}</font></p>
<table border="1" cellspacing="0" cellpadding="4" width="700">
{_html_rows(rows)}
</table>
<br>
{_html_grid(education_rows(p, truth.at("table")))}
<br>
{_html_grid(career_rows(p, truth.at("table")))}
<div class="memo">{_h(memo)}</div>
<h3>자기소개서</h3>
{letter}
<p>취미: 똠양꿍 만들기</p>
</body>
</html>
"""
    return html.encode("cp949"), {}


def html_r16(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("제로폭 공백으로 쪼갠 이름", "전각 숫자 전화번호", "인라인 흰 글씨", "font-size:0 숨은 텍스트", "HTML 주석·스크립트 속 이름(추출하지 않음)",
               "작성자 메타데이터", "증명사진")  # fmt: skip
    truth.put("metadata", "candidate", p.name)
    truth.put("body", "candidate", p.name)
    truth.put("body", "candidate", p.english)
    truth.put("body", "phone", phone_as(p, "fullwidth"))
    hidden = truth.at("hidden")
    white = f"{hidden('candidate', p.name)} {hidden('phone', p.phone_digits)} {hidden('email', p.email)}"
    zero = f"{p.name} 생년월일 {hidden('birthdate', birth_as(p, 'dots'))}"
    rows = profile_rows(p, truth.at("table"), Layout(phone="dash", birth="dots"))
    letter = "\n".join(f"<p>{_h(x)}</p>" for x in intro(p, truth, "body"))
    truth.images += 1
    spaced = "​".join(p.name)
    html = f"""<!doctype html>
<html lang="ko">
<head><meta charset="utf-8"><title>이력서</title><meta name="author" content="{p.name}"></head>
<body>
<!-- 작성자 {p.name}, 연락처 {p.phone} -->
<h1>이력서</h1>
<p>성명: {spaced} / {_h(p.english)}</p>
<p>휴대전화: {phone_as(p, 'fullwidth')}</p>
<p>키워드: 제어, PLC<span style="color:#ffffff">{_h(white)}</span></p>
<span style="font-size:0">{_h(zero)}</span>
<img src="{_photo_data_uri()}" alt="{p.name} 사진">
<table border="1">
{_html_rows(rows)}
</table>
{_html_grid(education_rows(p, truth.at("table")))}
{_html_grid(career_rows(p, truth.at("table")))}
<h2>자기소개서</h2>
{letter}
<script>var applicant = "{p.name}"; var phone = "{p.phone}";</script>
</body>
</html>
"""
    return html.encode("utf-8"), {}


def html_r17(p: Person, truth: Truth, font: Path) -> tuple[bytes, dict]:
    truth.put("filename", "candidate", p.name)
    truth.trap("Word에서 내보낸 HTML(mso 서식)", "인라인 태그로 쪼갠 이름(하<span>예린</span>)", "<wbr>로 쪼갠 이메일", "본문에 섞인 Word 머리글·바닥글",
               "조건부 주석 속 작성자(추출하지 않음)")  # fmt: skip
    truth.put("metadata", "candidate", p.name)
    body = truth.at("body")
    body("candidate", p.name)
    body("email", p.email)
    body("phone", p.phone)
    address_pieces(p, body)
    truth.put("link", "email", p.email)
    local, domain = p.email.split("@")
    rows = profile_rows(p, truth.at("table"), Layout(phone="spaces", birth="korean"))
    letter = "\n".join(f"<p class=MsoNormal>{_h(x)}<o:p></o:p></p>" for x in intro(p, truth, "body"))
    html = f"""<html xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"
xmlns:w="urn:schemas-microsoft-com:office:word" xmlns="http://www.w3.org/TR/REC-html40">
<head>
<meta http-equiv=Content-Type content="text/html; charset=utf-8">
<meta name=ProgId content=Word.Document>
<meta name=Generator content="Microsoft Word 15">
<title>{p.name} 이력서</title>
<!--[if gte mso 9]><xml>
 <o:DocumentProperties>
  <o:Author>{p.name}</o:Author>
  <o:LastAuthor>{p.english}</o:LastAuthor>
 </o:DocumentProperties>
</xml><![endif]-->
<style>
p.MsoNormal {{ margin: 0cm; font-size: 10.0pt; font-family: "맑은 고딕"; }}
@page WordSection1 {{ mso-header: h1; mso-footer: f1; }}
</style>
</head>
<body lang=KO style='tab-interval:40.0pt'>
<div class=WordSection1>
<p class=MsoTitle>이력서<o:p></o:p></p>
<p class=MsoNormal>성명: {p.name[0]}<span lang=EN-US style='font-family:"맑은 고딕"'>{p.name[1:]}</span><o:p></o:p></p>
<p class=MsoNormal>Email: <a href="mailto:{p.email}">{local}@<wbr>{domain}</a><o:p></o:p></p>
<p class=MsoNormal>주소: {p.region}<br>{p.road}<o:p></o:p></p>
<table class=MsoTableGrid border=1 cellspacing=0 cellpadding=0>
{_html_rows(rows)}
</table>
{_html_grid(education_rows(p, truth.at("table")))}
{_html_grid(career_rows(p, truth.at("table")))}
<p class=MsoNormal><b>자기소개서</b><o:p></o:p></p>
{letter}
<div style='mso-element:header' id=h1><p class=MsoHeader>{p.name} | {_h(p.job)} 지원 | {p.phone}<o:p></o:p></p></div>
<div style='mso-element:footer' id=f1><p class=MsoFooter>{p.name} 이력서 - 1 -<o:p></o:p></p></div>
</div>
</body>
</html>
"""
    return html.encode("utf-8"), {}


BUILDERS = {
    "R01": pdf_r01, "R02": pdf_r02, "R03": pdf_r03, "R04": pdf_r04, "R05": pdf_r05, "R06": pdf_r06,
    "R07": docx_r07, "R08": docx_r08, "R09": docx_r09, "R10": docx_r10,
    "R11": pptx_r11, "R12": pptx_r12, "R13": pptx_r13,
    "R14": html_r14, "R15": html_r15, "R16": html_r16, "R17": html_r17,
    "R18": scan_r18, "R19": scan_r19, "R20": scan_r20,
}  # fmt: skip



# --- main -----------------------------------------------------------------------------------


def find_font(arg: str | None) -> Path:
    for candidate in (arg, os.environ.get("IA_FIXTURE_FONT"), *FONT_CANDIDATES):
        if candidate and Path(candidate).is_file():
            return Path(candidate)
    raise SystemExit(
        "Korean TrueType font not found. Pass --font PATH or set IA_FIXTURE_FONT "
        "(e.g. malgun.ttf, NanumGothic.ttf). Fonts are read from the system, never committed."
    )


def build_all(font: Path, out_dir: Path) -> list[dict]:
    register_font(font)
    out_dir.mkdir(parents=True, exist_ok=True)
    entries = []
    for person in PEOPLE:
        truth = Truth()
        data, extra = BUILDERS[person.id](person, truth, font)
        (out_dir / person.file).write_bytes(data)
        entry: dict = {"id": person.id, "file": person.file}
        if extra.get("upload_name"):
            entry["upload_name"] = extra["upload_name"]
        entry["format"] = "pdf" if person.form == "scan" else person.form
        entry["scanned"] = person.form == "scan"
        if extra.get("ocr_pages"):
            entry["ocr_pages"] = extra["ocr_pages"]
        entry["traps"] = truth.traps
        entry["identity"] = person.identity()
        entry["expected"] = [
            {"category": c, "value": v, "where": kinds} for (c, v), kinds in truth.where.items()
        ]
        entry["keep"] = truth.keep
        entry["images"] = truth.images
        entries.append(entry)
    return entries


class _Flow(dict):
    """Dumped on one line: `- {category: phone, value: ..., where: [table]}`."""


class _Dumper(yaml.SafeDumper):
    pass


_Dumper.add_representer(_Flow, lambda d, data: d.represent_mapping("tag:yaml.org,2002:map", data, flow_style=True))


def write_manifest(entries: list[dict], path: Path) -> None:
    for entry in entries:
        entry["expected"] = [_Flow(e) for e in entry["expected"]]
    body = yaml.dump(
        {"samples": entries}, Dumper=_Dumper, allow_unicode=True, sort_keys=False, width=200, default_flow_style=None
    )
    path.write_text(MANIFEST_HEADER + body, encoding="utf-8", newline="\n")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--font", help="Korean TrueType font (.ttf); default: IA_FIXTURE_FONT or system paths")
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    font = find_font(args.font)
    entries = build_all(font, args.out)
    write_manifest(entries, args.out / "manifest.yaml")
    print(f"font {font}")
    for e in entries:
        print(f"{e['id']} {e['format']:<5} {'scan' if e['scanned'] else '    '} {len(e['expected']):>3} values  {e['file']}")
    print(f"wrote {len(entries)} files and manifest.yaml to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

