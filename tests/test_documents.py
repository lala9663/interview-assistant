"""Text extraction from files: every place text can hide is read, images are dropped, scans are
flagged, nothing is made up. Runs on the generated resumes in tests/fixtures/resumes."""

import io
import unicodedata
import zipfile

import pytest

from interview_assistant.deid.leaks import leak_pattern, load_document_samples
from interview_assistant.documents import (
    BlockKind,
    DocumentExtractor,
    ExtractedDocument,
    ExtractionError,
    UnsupportedFormat,
    detect_format,
    extract_document,
    normalize_text,
)
from interview_assistant.documents.base import BlockCollector
from interview_assistant.documents.pdf import unwrap_cell
from interview_assistant.documents.webpage import decode_html

from .conftest import ROOT

MANIFEST = ROOT / "tests" / "fixtures" / "resumes" / "manifest.yaml"


@pytest.fixture(scope="module")
def samples():
    return {s.id: s for s in load_document_samples(MANIFEST)}


@pytest.fixture(scope="module")
def docs(samples):
    return {i: extract_document(s.path.read_bytes(), s.upload_name) for i, s in samples.items()}


def texts(doc: ExtractedDocument, *kinds: BlockKind) -> str:
    return "\n".join(b.text for b in doc.blocks if not kinds or b.kind in kinds)


def test_fixture_set_matches_the_plan(samples):
    forms = [("scan" if s.scanned else s.format) for s in samples.values()]
    assert len(samples) == 20
    assert {f: forms.count(f) for f in set(forms)} == {"pdf": 6, "docx": 4, "pptx": 3, "html": 4, "scan": 3}
    assert all(s.expected and s.traps for s in samples.values())


def test_format_is_detected_from_content(samples, docs):
    for i, s in samples.items():
        assert docs[i].format == s.format, i
    docx_bytes = samples["R07"].path.read_bytes()
    assert detect_format(docx_bytes, "renamed.pdf") == "docx"


@pytest.mark.parametrize(
    "data, name, message",
    [
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, "이력서.hwp", "HWP"),
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, "old.doc", "HWP 또는 구형"),
        (b"GIF89a....", "photo.gif", "지원하지 않는"),
    ],
)
def test_unsupported_formats_are_refused(data, name, message):
    with pytest.raises(UnsupportedFormat, match=message):
        extract_document(data, name)


def test_hwpx_is_refused_by_name():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("mimetype", "application/hwp+zip")
        zf.writestr("Contents/section0.xml", "<x/>")
    with pytest.raises(UnsupportedFormat, match="HWPX"):
        extract_document(buf.getvalue(), "이력서.hwpx")


def test_broken_files_raise_extraction_error():
    with pytest.raises(ExtractionError):
        extract_document(b"%PDF-1.7\n garbage", "broken.pdf")
    with pytest.raises(ExtractionError):
        extract_document(b"PK\x03\x04 not really a zip", "broken.docx")


def test_every_planted_value_is_extracted_from_where_it_was_planted(samples, docs):
    """Header, footer, notes, link target, properties, file name, hidden text: if a value is in
    the file, extraction returns it, in a block of the kind it was put in."""
    gaps = []
    for i, s in samples.items():
        if s.scanned:
            continue
        blocks = [(b.kind.value, normalize_text(b.text)) for b in docs[i].blocks]
        for e in s.expected:
            pattern = leak_pattern(e)
            for kind in e.where:
                if not any(pattern.search(t) for k, t in blocks if k == kind):
                    gaps.append((i, e.value, kind))
    assert gaps == []


def test_kinds_of_blocks_each_format_yields(docs):
    def kinds(i):
        return {b.kind for b in docs[i].blocks}

    assert {BlockKind.filename, BlockKind.metadata, BlockKind.table, BlockKind.hidden, BlockKind.link} <= kinds("R01")
    assert BlockKind.comment in kinds("R02")  # PDF sticky note
    assert {BlockKind.header, BlockKind.footer, BlockKind.hidden, BlockKind.link} <= kinds("R07")
    assert BlockKind.comment in kinds("R08")  # Word review comment
    assert {BlockKind.notes, BlockKind.footer, BlockKind.hidden, BlockKind.link} <= kinds("R11")
    assert {BlockKind.header, BlockKind.footer, BlockKind.link} <= kinds("R14")


def test_properties_include_app_and_custom_fields(docs):
    assert docs["R07"].metadata["회사"] == "네이버"  # docProps/app.xml
    custom = [b.text for b in docs["R10"].blocks_of(BlockKind.metadata) if b.label == "기타 속성"]
    assert "지원자: 문태오" in custom and "연락처: 010-1450-6627" in custom  # docProps/custom.xml
    assert docs["R12"].metadata["마지막 수정자"] == "Minhyuk Pyo"
    assert docs["R14"].metadata["작성자"] == "육지안"  # <meta name="author">


def test_software_properties_are_not_extracted(docs):
    """'Microsoft Word 15' says nothing about the applicant and would become a company token."""
    for doc in docs.values():
        assert not any("Microsoft" in b.text or "ReportLab" in b.text for b in doc.blocks_of(BlockKind.metadata))


def test_images_are_counted_not_extracted(samples, docs):
    for i, s in samples.items():
        assert docs[i].images_dropped >= s.images, i
    rendered = docs["R14"].render() + docs["R16"].render()
    assert "data:image" not in rendered
    assert "증명사진" not in rendered and "사진" not in texts(docs["R16"])  # alt text is dropped too


def test_hidden_text_is_separated_from_visible_text(docs):
    assert "백엔드 우대" in texts(docs["R01"], BlockKind.hidden)  # white on white
    assert "우대" not in texts(docs["R01"], BlockKind.body)
    assert "추천인 이서준 선배" in texts(docs["R05"], BlockKind.hidden)  # with its spaces
    assert "eunchae.hb@example.org" in texts(docs["R05"], BlockKind.hidden)  # 0.5 pt text
    assert "예전 주소" in texts(docs["R08"], BlockKind.hidden)  # w:vanish
    assert "개인 연락처" in texts(docs["R12"], BlockKind.hidden)  # off the slide
    assert "비공개 메모" in texts(docs["R15"], BlockKind.hidden)  # CSS class display:none
    hidden16 = texts(docs["R16"], BlockKind.hidden)
    assert "01016623018" in hidden16 and "1997.11.28" in hidden16  # color:#fff, font-size:0


def test_white_text_on_a_dark_background_is_visible(docs):
    assert not docs["R03"].blocks_of(BlockKind.hidden)  # PDF sidebar
    assert "010.0447.1290" in texts(docs["R03"], BlockKind.body)
    assert "입사지원서 - 인프라 운영" in texts(docs["R15"], BlockKind.body)  # CSS banner
    assert "입사지원서" not in texts(docs["R15"], BlockKind.hidden)


def test_comments_scripts_and_conditional_comments_are_not_extracted(docs):
    assert "var applicant" not in docs["R16"].render()
    assert "작성자 마수빈" not in docs["R16"].render()  # <!-- ... -->
    assert "마지막 수정자" not in docs["R17"].metadata  # o:LastAuthor inside <!--[if gte mso 9]>


def test_file_specific_text_shapes_come_out_readable(docs):
    assert "마수빈" in texts(docs["R16"], BlockKind.body)  # zero-width spaces removed
    assert "010-1662-3018" in texts(docs["R16"], BlockKind.body)  # full-width digits folded
    assert "똠양꿍" in texts(docs["R15"])  # EUC-KR page decoded as CP949
    assert docs["R05"].filename == "황보은채 이력서.pdf"  # NFD upload name composed
    assert "하예린" in texts(docs["R17"], BlockKind.body)  # 하<span>예린</span>
    assert "yerin.ha@example.org" in texts(docs["R17"], BlockKind.body)  # <wbr> inside
    assert "haram-sunwoo.example.net" in texts(docs["R09"], BlockKind.table)  # nested table


def test_pdf_tables_keep_wrapped_cells_together(docs):
    profile = texts(docs["R06"], BlockKind.table)
    assert "junho.cha@example.net" in profile  # wrapped e-mail rejoined
    assert "차 준 호" in profile  # one syllable per line, kept as one cell


@pytest.mark.parametrize(
    "cell, expected",
    [
        ("harin.kwon@exa\nmple.com", "harin.kwon@example.com"),
        ("harin.kw\non@example.com", "harin.kwon@example.com"),
        ("harin.kwon\n@example.com", "harin.kwon@example.com"),
        ("010-0211-\n3344", "010-0211-3344"),
        ("https://github.com/\nharin-kwon", "https://github.com/harin-kwon"),
        ("010.0447.1290\nserin.oh@example.net", "010.0447.1290\nserin.oh@example.net"),
        ("serin.oh@example.net\nGitHub", "serin.oh@example.net\nGitHub"),
        ("서울특별시 마포구\n솔바람길 27", "서울특별시 마포구\n솔바람길 27"),
    ],
)
def test_unwrap_cell_only_glues_inside_addresses(cell, expected):
    assert unwrap_cell(cell) == expected


def test_scanned_pages_need_ocr_and_no_text_is_made_up(samples, docs):
    for i in ("R18", "R19", "R20"):
        doc = docs[i]
        assert doc.needs_ocr and doc.ocr_pages == list(samples[i].ocr_pages), i
        assert doc.images_dropped >= 1
    assert not docs["R18"].blocks_of(BlockKind.body, BlockKind.table)
    assert {b.text for b in docs["R19"].blocks_of(BlockKind.body)} == {"Scanned with ScanApp"}
    assert [b.label for b in docs["R20"].blocks_of(BlockKind.body)] == ["1쪽"]  # page 2 is the scan
    for i, s in samples.items():
        if not s.scanned:
            assert not docs[i].needs_ocr, i


def test_pictures_without_text_need_ocr_in_every_format():
    """A scan pasted into Word or a web page is still a scan."""
    import docx
    from docx.shared import Cm
    from PIL import Image

    png = io.BytesIO()
    Image.new("L", (40, 60), 200).save(png, format="PNG")
    document = docx.Document()
    document.add_paragraph().add_run().add_picture(io.BytesIO(png.getvalue()), width=Cm(15))
    buf = io.BytesIO()
    document.save(buf)
    doc = extract_document(buf.getvalue(), "scan.docx")
    assert doc.needs_ocr and doc.ocr_pages == [1]

    page = '<html><body><img src="scan.jpg" alt="김도윤 이력서"></body></html>'.encode()
    doc = extract_document(page, "scan.html")
    assert doc.needs_ocr and doc.images_dropped == 1


def test_normalize_text():
    assert normalize_text(unicodedata.normalize("NFD", "권하린")) == "권하린"
    assert normalize_text("마​수‌빈﻿") == "마수빈"
    assert normalize_text("０１０－１６６２") == "010-1662"
    assert normalize_text("제갈\xa0윤호　님\t\t입니다") == "제갈 윤호 님 입니다"
    assert normalize_text("a\r\nb\vc\n\n\n\nd") == "a\nb\nc\n\nd"


def test_decode_html_reads_legacy_korean():
    page = '<meta charset="euc-kr"><p>똠방각하 김도윤</p>'.encode("cp949")
    assert "똠방각하 김도윤" in decode_html(page)
    assert "김도윤" in decode_html("<p>김도윤</p>".encode("cp949"))  # no declaration


def test_unknown_metadata_keys_keep_their_name_in_the_text():
    out = BlockCollector("a.pdf", "pdf")
    out.add_metadata([("지원자 연락처", "010-0123-4567"), ("Author", "김도윤"), ("Created", "2026-01-01T00:00:00Z")])
    blocks = [(b.label, b.text) for b in out.build().blocks_of(BlockKind.metadata)]
    assert blocks == [("기타 속성", "지원자 연락처: 010-0123-4567"), ("작성자", "김도윤")]


def test_extractors_are_replaceable(samples):
    class Stub(DocumentExtractor):
        format = "pdf"

        def extract(self, data, filename):
            out = BlockCollector(filename, self.format)
            out.add(BlockKind.body, "stub", "from the stub")
            return out.build()

    doc = extract_document(samples["R01"].path.read_bytes(), "x.pdf", {"pdf": Stub()})
    assert texts(doc, BlockKind.body) == "from the stub"
    with pytest.raises(UnsupportedFormat):
        extract_document(samples["R07"].path.read_bytes(), "x.docx", {"pdf": Stub()})


def test_extraction_makes_no_network_calls(samples, no_network):
    for s in samples.values():
        extract_document(s.path.read_bytes(), s.upload_name)
