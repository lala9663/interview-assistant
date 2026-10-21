"""PPTX extractor on python-pptx (MIT).

Per slide: text frames (inside groups too), tables cell by cell, speaker notes, run and shape
hyperlinks. Text a viewer would not see in the slide show (white text on a shape without fill,
shapes placed outside the slide) goes to a hidden block. Text the slide master or layouts add
to every slide (a footer with the presenter's name, say) is read once, from the non-placeholder
shapes of each master and layout in use. A slide that is a picture with almost no text (a
scanned page pasted in) is marked as needing OCR, like a PDF page.
"""

from __future__ import annotations

import io
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

from interview_assistant.documents.base import BlockCollector, DocumentExtractor, ExtractionError
from interview_assistant.documents.normalize import collapse
from interview_assistant.documents.ooxml import count_images, read_properties
from interview_assistant.documents.types import BlockKind, ExtractedDocument, PageInfo

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_WHITE_SCHEMES = frozenset({"bg1", "lt1"})


class _Sink:
    def __init__(self) -> None:
        self.text: list[str] = []
        self.tables: list[str] = []
        self.hidden: list[str] = []
        self.links: list[str] = []
        self.picture_area = 0


class PptxExtractor(DocumentExtractor):
    format = "pptx"

    def __init__(self, *, min_text_chars: int = 30, scan_coverage: float = 0.5) -> None:
        self.min_text_chars = min_text_chars
        self.scan_coverage = scan_coverage

    def extract(self, data: bytes, filename: str) -> ExtractedDocument:
        try:
            prs = Presentation(io.BytesIO(data))
        except Exception as exc:
            raise ExtractionError(f"PPTX를 읽을 수 없습니다: {type(exc).__name__}: {exc}") from None
        out = BlockCollector(filename, self.format)
        out.add_metadata(read_properties(data))
        size = (int(prs.slide_width or 0), int(prs.slide_height or 0))

        templates: list[Any] = []
        links: list[str] = []
        for n, slide in enumerate(prs.slides, start=1):
            hidden_slide = slide._element.get("show") in ("0", "false")
            label = f"슬라이드 {n}" + (" (숨김)" if hidden_slide else "")
            sink = _Sink()
            _shapes(slide.shapes, sink, size, top_level=True)
            out.add(BlockKind.body, label, "\n".join(sink.text))
            for i, table in enumerate(sink.tables, start=1):
                out.add(BlockKind.table, f"{label} 표 {i}", table)
            out.add(BlockKind.hidden, f"{label} 숨은 텍스트", "\n".join(sink.hidden))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = _Sink()
                out.add(BlockKind.notes, f"{label} 노트", _frame_text(slide.notes_slide.notes_text_frame, notes, filled=True))
                links.extend(notes.links)
            links.extend(sink.links)

            text_chars = sum(len("".join(t.split())) for t in (*sink.text, *sink.tables))
            coverage = min(1.0, sink.picture_area / (size[0] * size[1])) if all(size) else 0.0
            needs_ocr = text_chars < self.min_text_chars and coverage >= self.scan_coverage
            out.pages.append(PageInfo(n, text_chars, round(coverage, 3), needs_ocr))
            if needs_ocr:
                out.warnings.append(f"{label}: 그림만 있고 텍스트 없음, OCR 필요")
            for template in (slide.slide_layout, slide.slide_layout.slide_master):
                if all(template is not t for t in templates):
                    templates.append(template)

        for template in templates:
            label = "슬라이드 마스터" if template.__class__.__name__ == "SlideMaster" else "슬라이드 레이아웃"
            sink = _Sink()
            _shapes(template.shapes, sink, size, top_level=True, skip_placeholders=True)
            out.add(BlockKind.footer, label, "\n".join([*sink.text, *sink.tables]))
            out.add(BlockKind.hidden, f"{label} 숨은 텍스트", "\n".join(sink.hidden))
            links.extend(sink.links)

        out.add_links("링크", links)
        out.images = count_images(data)
        return out.build()


def _shapes(
    shapes: Any,
    sink: _Sink,
    size: tuple[int, int],
    *,
    top_level: bool,
    skip_placeholders: bool = False,
) -> None:
    for shape in shapes:
        if skip_placeholders and shape.is_placeholder:
            continue
        off_slide = top_level and _off_slide(shape, size)
        target = _Sink() if off_slide else sink
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _shapes(shape.shapes, target, size, top_level=False)
        if top_level and not off_slide and _is_picture(shape):
            sink.picture_area += _area(shape, size)
        if shape.has_text_frame:
            target.text.append(_frame_text(shape.text_frame, target, filled=_has_fill(shape)))
        if getattr(shape, "has_table", False):
            rows = []
            for row in shape.table.rows:
                cells = [collapse(_frame_text(c.text_frame, target, filled=True)) for c in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            target.tables.append("\n".join(rows))
        try:
            address = shape.click_action.hyperlink.address
        except Exception:  # not every shape type supports click actions
            address = None
        if address:
            target.links.append(address)
        if off_slide:  # nothing outside the slide is shown: all of it is hidden text
            sink.hidden.extend(t for t in (*target.text, *target.tables, *target.hidden) if t)
            sink.links.extend(target.links)


def _frame_text(frame: Any, sink: _Sink, *, filled: bool) -> str:
    lines = []
    for paragraph in frame.paragraphs:
        visible: list[str] = []
        secret: list[str] = []
        for node in paragraph._p:
            if node.tag == A + "br":
                visible.append("\n")
            elif node.tag in (A + "r", A + "fld"):
                text = "".join(t.text or "" for t in node.iter(A + "t"))
                (secret if not filled and _white(node) else visible).append(text)
                link = node.find(f"{A}rPr/{A}hlinkClick")
                if link is not None and (address := _rel_target(paragraph, link.get(R_ID))):
                    sink.links.append(address)
        lines.append("".join(visible))
        if "".join(secret).strip():
            sink.hidden.append("".join(secret))
    return "\n".join(lines)


def _rel_target(paragraph: Any, rid: str | None) -> str | None:
    if not rid:
        return None
    try:
        rel = paragraph.part.rels[rid]
    except (KeyError, AttributeError):
        return None
    return rel.target_ref if rel.is_external else None


def _white(run: etree._Element) -> bool:
    fill = run.find(f"{A}rPr/{A}solidFill")
    if fill is None:
        return False
    rgb = fill.find(A + "srgbClr")
    if rgb is not None:
        return (rgb.get("val") or "").upper() == "FFFFFF"
    scheme = fill.find(A + "schemeClr")
    return scheme is not None and scheme.get("val") in _WHITE_SCHEMES


def _has_fill(shape: Any) -> bool:
    """White text on a filled shape (a coloured box) is visible."""
    try:
        return shape.fill.type not in (None, MSO_FILL.BACKGROUND)
    except Exception:
        return False


def _is_picture(shape: Any) -> bool:
    return shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image")


def _area(shape: Any, size: tuple[int, int]) -> int:
    width, height = size
    if None in (shape.left, shape.top, shape.width, shape.height):
        return 0
    x0, y0 = max(0, shape.left), max(0, shape.top)
    x1, y1 = min(width, shape.left + shape.width), min(height, shape.top + shape.height)
    return max(0, x1 - x0) * max(0, y1 - y0)


def _off_slide(shape: Any, size: tuple[int, int]) -> bool:
    width, height = size
    left, top, w, h = shape.left, shape.top, shape.width, shape.height
    if None in (left, top, w, h) or not width or not height:
        return False
    return left >= width or top >= height or left + w <= 0 or top + h <= 0
